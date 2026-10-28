"""Unit tests of the Word copy tool (run inside the openwebui container: python-docx).

    docker compose exec openwebui python3 -m unittest discover -s /seed/tools -v
"""
import asyncio
import io
import os
import sys
import tempfile
import types
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import word_copy  # noqa: E402


def sample():
    doc = Document()
    doc.add_heading("Note de cadrage", level=1)
    p = doc.add_paragraph()
    p.add_run("Budget :").bold = True
    p.add_run(" 120 k€ sur deux ans.")
    doc.add_paragraph("Les justificatifs sont à envoyer sous 30 jours.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Pilote : DSI"
    return doc


def texts(doc):
    return [p.text for p in word_copy.all_paragraphs(doc)]


class DocumentText(unittest.TestCase):
    def test_includes_headings_body_and_tables(self):
        text = word_copy.document_text(sample())
        for part in ("# Note de cadrage", "Budget : 120 k€", "30 jours", "Pilote : DSI"):
            self.assertIn(part, text)

    def test_truncation_is_announced(self):
        doc = Document()
        doc.add_paragraph("x" * (word_copy.MAX_TEXT + 10))
        self.assertIn("tronqué", word_copy.document_text(doc))


class CopyName(unittest.TestCase):
    def test_no_path_and_no_markdown_breakers(self):
        name = word_copy.copy_name("../../etc/[note](x).docx")
        self.assertEqual(name, "note x (modifié).docx")


class ApplyChanges(unittest.TestCase):
    def setUp(self):
        self.doc = sample()

    def apply(self, *changes):
        return word_copy.apply_changes(self.doc, list(changes))

    def test_replace_keeps_formatting_of_other_runs(self):
        done, failed = self.apply({"action": "remplacer", "cible": "120 k€", "texte": "150 k€"})
        self.assertEqual((done, failed), (1, []))
        runs = self.doc.paragraphs[1].runs
        self.assertEqual(self.doc.paragraphs[1].text, "Budget : 150 k€ sur deux ans.")
        self.assertTrue(runs[0].bold)

    def test_replace_spanning_runs_and_extra_spaces(self):
        done, failed = self.apply({"action": "remplacer", "cible": "Budget  : 120", "texte": "Enveloppe : 150"})
        self.assertEqual((done, failed), (1, []))
        paragraph = self.doc.paragraphs[1]
        self.assertEqual(paragraph.text, "Enveloppe : 150 k€ sur deux ans.")
        # The new text takes the formatting where the match starts; the rest keeps its own.
        self.assertEqual([(r.text, bool(r.bold)) for r in paragraph.runs if r.text],
                         [("Enveloppe : 150", True), (" k€ sur deux ans.", False)])

    def test_text_inside_a_hyperlink_is_reported_not_duplicated(self):
        paragraph = self.doc.add_paragraph("Voir ")
        link = OxmlElement("w:hyperlink")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "le site intranet"
        run.append(text)
        link.append(run)
        paragraph._p.append(link)
        done, failed = self.apply({"action": "remplacer", "cible": "site intranet", "texte": "portail"})
        self.assertEqual(done, 0)
        self.assertIn("lien", failed[0])
        self.assertEqual(paragraph.text, "Voir le site intranet")

    def test_delete_only_paragraph_of_a_cell_keeps_the_cell_valid(self):
        self.apply({"action": "supprimer", "cible": "Pilote : DSI"})
        cell = self.doc.tables[0].cell(0, 0)
        self.assertEqual(len(cell.paragraphs), 1)
        self.assertEqual(cell.text, "")

    def test_malformed_changes_are_reported(self):
        done, failed = self.apply("pas un objet", {"action": "remplacer", "cible": None, "texte": None})
        self.assertEqual(done, 0)
        self.assertEqual(len(failed), 2)

    def test_replace_in_table_cell(self):
        self.apply({"action": "remplacer", "cible": "DSI", "texte": "Direction des opérations"})
        self.assertIn("Pilote : Direction des opérations", texts(self.doc))

    def test_insert_after_heading_uses_body_style_and_keeps_order(self):
        self.apply({"action": "inserer_apres", "cible": "Note de cadrage", "texte": "Ligne A\nLigne B"})
        paragraphs = self.doc.paragraphs
        self.assertEqual([p.text for p in paragraphs[1:3]], ["Ligne A", "Ligne B"])
        self.assertEqual(paragraphs[1].style.name, "Normal")

    def test_bare_heading_marker_is_not_a_target(self):
        done, failed = self.apply({"action": "supprimer", "cible": "# "})
        self.assertEqual(done, 0)
        self.assertEqual(self.doc.paragraphs[0].text, "Note de cadrage")

    def test_literal_hash_in_text_can_be_targeted(self):
        self.doc.add_paragraph("#2 Décision reportée")
        done, _ = self.apply({"action": "remplacer", "cible": "#2 Décision", "texte": "#3 Décision"})
        self.assertEqual(done, 1)
        self.assertIn("#3 Décision reportée", texts(self.doc))

    def test_inserts_on_the_same_target_keep_their_order(self):
        self.apply({"action": "inserer_apres", "cible": "30 jours", "texte": "A"},
                   {"action": "inserer_apres", "cible": "30 jours", "texte": "B"},
                   {"action": "inserer_tableau", "cible": "30 jours", "lignes": [["T"]]})
        body = [el for el in self.doc.element.body.iterchildren() if el.tag.split("}")[1] in ("p", "tbl")]
        kinds = ["tbl" if el.tag.endswith("tbl") else "".join(t.text or "" for t in el.iter(qn("w:t"))) for el in body]
        start = kinds.index("Les justificatifs sont à envoyer sous 30 jours.")
        self.assertEqual(kinds[start + 1:start + 4], ["A", "B", "tbl"])

    def test_too_many_changes_are_refused(self):
        done, failed = self.apply(*[{"action": "supprimer", "cible": "zzz"}] * (word_copy.MAX_CHANGES + 1))
        self.assertEqual(done, 0)
        self.assertEqual(len(failed), 1)

    def test_heading_marker_from_lire_document_is_accepted_in_target(self):
        done, failed = self.apply({"action": "inserer_apres", "cible": "# Note de cadrage", "texte": "Intro"})
        self.assertEqual((done, failed), (1, []))
        self.assertEqual(self.doc.paragraphs[1].text, "Intro")

    def test_delete_paragraph(self):
        self.apply({"action": "supprimer", "cible": "30 jours"})
        self.assertNotIn("Les justificatifs sont à envoyer sous 30 jours.", texts(self.doc))

    def test_failures_are_reported_and_others_applied(self):
        done, failed = self.apply(
            {"action": "remplacer", "cible": "introuvable", "texte": "x"},
            {"action": "renommer", "cible": "Budget", "texte": "x"},
            {"action": "supprimer", "cible": "30 jours"})
        self.assertEqual(done, 1)
        self.assertEqual(len(failed), 2)
        self.assertIn("introuvable", failed[0])
        self.assertIn("renommer", failed[1])


def reloaded(doc):
    """Round trip through a file: Word-level validity, not just in-memory XML."""
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return Document(buffer)


class Visuals(unittest.TestCase):
    ROWS = [["Étape", "Livrable"], ["POC", "Rapport de test"], ["Recette", "PV signé"]]
    STEPS = ["Kickoff", "Cadrage détaillé des besoins métier et des contraintes", "POC", "Recette", "Déploiement"]

    def setUp(self):
        self.doc = sample()

    def test_table_inserted_right_after_target_with_bold_header(self):
        done, failed = word_copy.apply_changes(self.doc, [
            {"action": "inserer_tableau", "cible": "Budget", "lignes": self.ROWS}])
        self.assertEqual((done, failed), (1, []))
        doc = reloaded(self.doc)
        anchor = doc.paragraphs[1]._p
        self.assertEqual(anchor.getnext().tag.split("}")[1], "tbl")
        table = doc.tables[0]
        self.assertEqual([c.text for c in table.rows[1].cells], ["POC", "Rapport de test"])
        self.assertTrue(all(r.bold for r in table.rows[0].cells[0].paragraphs[0].runs))

    def test_table_borders_respect_word_schema_order(self):
        word_copy.apply_changes(self.doc, [{"action": "inserer_tableau", "cible": "Budget", "lignes": self.ROWS}])
        children = [c.tag.split("}")[1] for c in self.doc.paragraphs[1]._p.getnext().tblPr]
        self.assertLess(children.index("tblBorders"), children.index("tblLook"))

    def test_none_cell_is_empty_and_nested_cell_is_refused(self):
        done, _ = word_copy.apply_changes(self.doc, [
            {"action": "inserer_tableau", "cible": "Budget", "lignes": [["a", None]]}])
        self.assertEqual(done, 1)
        self.assertEqual(self.doc.paragraphs[1]._p.getnext().xpath("string(.)"), "a")
        done, failed = word_copy.apply_changes(self.doc, [
            {"action": "inserer_tableau", "cible": "Budget", "lignes": [["a", ["b"]]]},
            {"action": "inserer_tableau", "cible": "Budget", "lignes": [["x" * (word_copy.MAX_CELL + 1)]]},
            {"action": "inserer_schema_etapes", "cible": "Budget", "etapes": ["a", "x" * (word_copy.MAX_LABEL + 1)]}])
        self.assertEqual((done, len(failed)), (0, 3))

    def test_diagram_has_alt_text_and_readable_labels(self):
        word_copy.apply_changes(self.doc, [{"action": "inserer_schema_etapes", "cible": "30 jours", "etapes": self.STEPS}])
        self.assertIn("Kickoff", self.doc.inline_shapes[0]._inline.docPr.get("descr"))
        points = word_copy.FONT_SIZE * (16 / 2.54 * 72) / word_copy.full_row_width()
        self.assertGreaterEqual(points, 9)

    def test_table_in_a_cell_keeps_the_cell_ending_with_a_paragraph(self):
        word_copy.apply_changes(self.doc, [{"action": "inserer_tableau", "cible": "Pilote", "lignes": self.ROWS}])
        cell = reloaded(self.doc).tables[0].cell(0, 0)
        self.assertEqual(cell._tc[-1].tag.split("}")[1], "p")

    def test_step_diagram_is_a_png_picture_after_target(self):
        done, failed = word_copy.apply_changes(self.doc, [
            {"action": "inserer_schema_etapes", "cible": "30 jours", "etapes": self.STEPS}])
        self.assertEqual((done, failed), (1, []))
        doc = reloaded(self.doc)
        self.assertEqual(len(doc.inline_shapes), 1)
        texts = [p.text for p in doc.paragraphs]
        picture = next(i for i, p in enumerate(doc.paragraphs) if p._p.xpath(".//w:drawing"))
        self.assertEqual(texts[picture - 1], "Les justificatifs sont à envoyer sous 30 jours.")

    def test_wrapped_labels_fit_their_box(self):
        font = word_copy.load_font()
        lines = word_copy.wrap(" ".join([self.STEPS[1]] * 3), font, 330)
        self.assertEqual(len(lines), 3)
        self.assertTrue(all(font.getlength(line) <= 330 for line in lines), lines)

    def test_draw_steps_returns_png(self):
        self.assertTrue(word_copy.draw_steps(self.STEPS).startswith(b"\x89PNG"))

    def test_invalid_visual_arguments_are_reported(self):
        done, failed = word_copy.apply_changes(self.doc, [
            {"action": "inserer_tableau", "cible": "Budget", "lignes": []},
            {"action": "inserer_tableau", "cible": "Budget", "lignes": [["a"], "pas une ligne"]},
            {"action": "inserer_schema_etapes", "cible": "Budget", "etapes": ["seule"]},
            {"action": "inserer_schema_etapes", "cible": "Budget", "etapes": ["x"] * 13}])
        self.assertEqual(done, 0)
        self.assertEqual(len(failed), 4)


class PickFile(unittest.TestCase):
    FILES = [{"type": "file", "id": "a", "name": "cadrage.docx"},
             {"type": "file", "id": "b", "name": "budget.docx"},
             {"type": "file", "id": "c", "name": "notes.pdf"}]

    def test_by_name(self):
        self.assertEqual(word_copy.pick_file(self.FILES, "budget.docx"), "b")

    def test_single_docx_needs_no_name(self):
        self.assertEqual(word_copy.pick_file(self.FILES[:1] + self.FILES[2:], ""), "a")

    def test_ambiguous_or_missing_is_refused(self):
        with self.assertRaises(ValueError):
            word_copy.pick_file(self.FILES, "")
        with self.assertRaises(ValueError):
            word_copy.pick_file(self.FILES[2:], "")


class Ownership(unittest.TestCase):
    def test_only_the_owner_may_read(self):
        word_copy.check_owner("u1", {"id": "u1"})
        with self.assertRaises(PermissionError):
            word_copy.check_owner("u2", {"id": "u1"})


class LoadAttached(unittest.TestCase):
    """load_attached with fake Open WebUI modules: no database needed."""

    def setUp(self):
        self.path = os.path.join(tempfile.mkdtemp(), "f.docx")
        sample().save(self.path)
        self.read = []
        self.record = types.SimpleNamespace(user_id="owner", filename="cadrage.docx", path=self.path)
        record, read = self.record, self.read

        class Files:
            @staticmethod
            async def get_file_by_id(file_id):
                return record

        class Storage:
            @staticmethod
            def get_file(path):
                read.append(path)
                return path

        sys.modules["open_webui.models.files"] = types.SimpleNamespace(Files=Files)
        sys.modules["open_webui.storage.provider"] = types.SimpleNamespace(Storage=Storage)
        self.addCleanup(sys.modules.pop, "open_webui.models.files", None)
        self.addCleanup(sys.modules.pop, "open_webui.storage.provider", None)

    def read_as(self, user_id):
        files = [{"type": "file", "id": "a", "name": "cadrage.docx"}]
        return asyncio.run(word_copy.Tools().lire_document_word("", {"id": user_id}, files))

    def test_owner_reads_the_whole_document(self):
        self.assertIn("Pilote : DSI", self.read_as("owner"))

    def test_foreign_file_is_refused_before_any_read(self):
        self.assertIn("n'appartient pas", self.read_as("intruder"))
        self.assertEqual(self.read, [])

    def test_stored_file_that_is_not_a_docx_is_refused(self):
        self.record.filename = "notes.pdf"
        self.assertIn(".docx", self.read_as("owner"))
        self.assertEqual(self.read, [])

    def test_corrupt_docx_is_reported(self):
        with open(self.path, "w") as f:
            f.write("not a zip")
        self.assertIn("illisible", self.read_as("owner"))


if __name__ == "__main__":
    unittest.main()
