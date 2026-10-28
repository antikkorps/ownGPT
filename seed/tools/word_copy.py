"""
title: Copie modifiée d'un document Word
description: Lit en entier un .docx joint au chat et en produit une copie modifiée à télécharger (l'original n'est jamais touché).
version: 1.0.0
"""
import asyncio
import io
import os
import re
import uuid

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

# Upper bound on the text handed to the model, so a huge file cannot flood the context.
MAX_TEXT = 60000


def norm(text):
    return " ".join(text.split())


def all_paragraphs(doc):
    """Body paragraphs, then table cells: a cadrage note often keeps key facts in tables."""
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def is_heading(paragraph):
    return paragraph.style is not None and paragraph.style.name.lower().startswith(("heading", "titre", "title"))


def document_text(doc):
    lines = [("# " if is_heading(p) else "") + p.text for p in all_paragraphs(doc) if p.text.strip()]
    text = "\n".join(lines)
    if len(text) > MAX_TEXT:
        return text[:MAX_TEXT] + f"\n[… texte tronqué à {MAX_TEXT} caractères]"
    return text


def find_paragraph(doc, target):
    wanted = norm(target)
    return next((p for p in all_paragraphs(doc) if wanted in norm(p.text)), None)


def replace_text(paragraph, old, new):
    """Splices the match out of the runs it covers: the new text takes the formatting where
    the match starts, text around it keeps its own. Spacing may differ from the document."""
    runs = paragraph.runs
    pattern = r"\s+".join(map(re.escape, old.split()))
    match = re.search(pattern, "".join(r.text for r in runs))
    if not match:
        # paragraph.text also covers hyperlinks and fields, which paragraph.runs does not.
        raise ValueError(f"« {old[:80]} » est dans un lien ou un champ : non modifiable")
    start, end = match.span()
    pos, first = 0, True
    for run in runs:
        run_start, pos = pos, pos + len(run.text)
        if pos <= start or run_start >= end:
            continue
        head = run.text[:max(start - run_start, 0)] if first else ""
        tail = run.text[end - run_start:] if end < pos else ""
        run.text = head + (new if first else "") + tail
        first = False


def insert_after(paragraph, text):
    # Text inserted under a heading is body text, not another heading.
    style = None if is_heading(paragraph) else paragraph.style
    anchor = paragraph
    for line in text.split("\n"):
        element = OxmlElement("w:p")
        anchor._p.addnext(element)
        anchor = Paragraph(element, paragraph._parent)
        if style is not None:
            anchor.style = style
        anchor.add_run(line)


def delete_paragraph(paragraph):
    p = paragraph._p
    parent = p.getparent()
    # A table cell needs at least one paragraph and a section break lives in its paragraph:
    # removing either corrupts the file, so those are emptied instead.
    if (parent.tag == qn("w:tc") and len(parent.findall(qn("w:p"))) == 1) or p.find(".//" + qn("w:sectPr")) is not None:
        for child in list(p):
            if child.tag != qn("w:pPr"):
                p.remove(child)
        return
    parent.remove(p)


ACTIONS = {
    "remplacer": lambda p, c: replace_text(p, c["cible"], c.get("texte") or ""),
    "inserer_apres": lambda p, c: insert_after(p, c.get("texte") or ""),
    "supprimer": lambda p, c: delete_paragraph(p),
}


def invalid(change):
    """Why a change from the model cannot be applied, or None."""
    if not isinstance(change, dict):
        return "modification ignorée : objet {action, cible, texte} attendu"
    if change.get("action") not in ACTIONS:
        return f"action inconnue « {change.get('action')} » (remplacer, inserer_apres, supprimer)"
    if not isinstance(change.get("cible"), str) or not change["cible"].strip():
        return f"{change['action']} : « cible » manquante"
    if not isinstance(change.get("texte") or "", str):
        return f"{change['action']} : « texte » doit être du texte"
    return None


def apply_changes(doc, changes):
    """Applies what it can; returns (applied count, failure messages) so the model can retry."""
    done, failed = 0, []
    for change in changes:
        reason = invalid(change)
        paragraph = None if reason else find_paragraph(doc, change["cible"])
        if not reason and paragraph is None:
            reason = f"{change['action']} : texte introuvable « {change['cible'][:80]} »"
        if not reason:
            try:
                ACTIONS[change["action"]](paragraph, change)
                done += 1
            except ValueError as e:
                reason = str(e)
        if reason:
            failed.append(reason)
    return done, failed


def pick_file(files, name):
    docx = [f for f in files or [] if f.get("type") == "file" and f.get("name", "").lower().endswith(".docx")]
    if name:
        docx = [f for f in docx if f.get("name") == name]
    if len(docx) != 1:
        found = ", ".join(f.get("name", "?") for f in docx) or "aucun"
        raise ValueError(f"Préciser un seul fichier .docx joint au chat (trouvé : {found}).")
    return docx[0]["id"]


def check_owner(file_user_id, user):
    # A chat can only reference files, but the id comes from the request: never trust it alone.
    if file_user_id != (user or {}).get("id"):
        raise PermissionError("Ce fichier n'appartient pas à l'utilisateur.")


def copy_name(filename):
    # Upload names are user-chosen: no path parts, nothing that would break the markdown link.
    stem = os.path.splitext(os.path.basename(filename or ""))[0]
    stem = norm(re.sub(r"[\[\]()<>]", " ", stem))
    return f"{stem or 'document'} (modifié).docx"


async def load_attached(files, name, user):
    # Open WebUI internals are imported here so the helpers above stay testable anywhere.
    from open_webui.models.files import Files
    from open_webui.storage.provider import Storage

    record = await Files.get_file_by_id(pick_file(files, name))
    if record is None:
        raise ValueError("Fichier introuvable.")
    check_owner(record.user_id, user)
    if not (record.filename or "").lower().endswith(".docx"):
        raise ValueError("Seuls les fichiers Word .docx sont pris en charge.")
    path = await asyncio.to_thread(Storage.get_file, record.path)
    try:
        return record.filename, Document(path)
    except Exception:
        raise ValueError(f"Document Word illisible : {record.filename}.") from None


async def store_copy(doc, filename, user):
    from open_webui.models.files import FileForm, Files
    from open_webui.storage.provider import Storage

    buffer = io.BytesIO()
    doc.save(buffer)
    file_id, name = str(uuid.uuid4()), copy_name(filename)
    buffer.seek(0)
    tags = {"OpenWebUI-User-Id": user["id"], "OpenWebUI-File-Id": file_id}
    _, path = await asyncio.to_thread(Storage.upload_file, buffer, f"{file_id}_{name}", tags)
    await Files.insert_new_file(user["id"], FileForm(id=file_id, filename=name, path=path, meta={
        "name": name, "size": buffer.getbuffer().nbytes,
        "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}))
    return file_id, name


class Tools:
    async def lire_document_word(self, nom_fichier: str = "", __user__: dict = None, __files__: list = None) -> str:
        """
        Renvoie le texte intégral d'un document Word (.docx) joint au chat, titres marqués par « # ».
        À appeler avant de proposer ou d'appliquer des modifications : sans cet outil, seuls des extraits du document sont visibles.
        :param nom_fichier: Nom du fichier .docx ; vide s'il n'y en a qu'un dans le chat.
        """
        try:
            filename, doc = await load_attached(__files__, nom_fichier, __user__)
        except (ValueError, PermissionError) as e:
            return str(e)
        return f"Contenu intégral de {filename} :\n\n{document_text(doc)}"

    async def modifier_copie_word(self, modifications: list[dict], nom_fichier: str = "", __user__: dict = None, __files__: list = None) -> str:
        """
        Crée une copie modifiée d'un document Word (.docx) joint au chat, avec sa mise en forme ; l'original n'est pas modifié.
        Utiliser cet outil dès que l'utilisateur demande de modifier, corriger ou compléter le document, au lieu de répondre qu'il est impossible de modifier un fichier.
        :param modifications: Liste d'objets {"action": "remplacer" | "inserer_apres" | "supprimer", "cible": texte exact présent dans le document (un paragraphe ou une partie), "texte": nouveau texte (remplacer, inserer_apres ; une ligne par paragraphe)}.
        :param nom_fichier: Nom du fichier .docx ; vide s'il n'y en a qu'un dans le chat.
        """
        if not isinstance(modifications, list):
            return "« modifications » doit être une liste d'objets {action, cible, texte}."
        try:
            filename, doc = await load_attached(__files__, nom_fichier, __user__)
        except (ValueError, PermissionError) as e:
            return str(e)
        done, failed = apply_changes(doc, modifications)
        if not done:
            return "Aucune modification appliquée : " + " ; ".join(failed)
        file_id, name = await store_copy(doc, filename, __user__)
        report = f" Non appliquées : {' ; '.join(failed)}." if failed else ""
        return (f"Copie créée avec {done} modification(s).{report} Donne ce lien à l'utilisateur tel quel : "
                f"[Télécharger {name}](/api/v1/files/{file_id}/content?attachment=true)")
